"""Excel 合并表头 / 标题行解析回归。"""
from pathlib import Path

import pandas as pd
import pytest
from openpyxl import Workbook

from service.data_ingestion.readers.excel_reader import ExcelReader, _merged_ranges


def _new(path: Path):
    wb = Workbook()
    return wb, wb.active


def _two_level_header(ws, start_row: int = 1):
    """写一个两行合并表头（A 列纵向合并 + 年度横向合并）。"""
    r = start_row
    ws.cell(r, 1, "城市")
    ws.merge_cells(start_row=r, start_column=1, end_row=r + 1, end_column=1)
    ws.cell(r, 2, "2023 年")
    ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=3)
    ws.cell(r, 4, "2024 年")
    ws.merge_cells(start_row=r, start_column=4, end_row=r, end_column=5)
    for col, (year, metric) in enumerate(
        [(None, None), ("2023", "销售额"), ("2023", "利润"), ("2024", "销售额"), ("2024", "利润")],
        start=1,
    ):
        if metric:
            ws.cell(r + 1, col, metric)
    data = [
        ["北京", 1200, 300, 1500, 400],
        ["上海", 1100, 280, 1400, 360],
        ["广州", 900, 210, 1250, 300],
    ]
    for offset, row in enumerate(data, start=r + 2):
        for col, value in enumerate(row, start=1):
            ws.cell(offset, col, value)


def test_two_level_merged_header(tmp_path: Path):
    path = tmp_path / "sales.xlsx"
    wb, ws = _new(path)
    _two_level_header(ws)
    wb.save(path)

    table = ExcelReader.read(str(path))[0]

    assert list(table.df.columns) == [
        "城市",
        "2023 年_销售额",
        "2023 年_利润",
        "2024 年_销售额",
        "2024 年_利润",
    ]
    assert len(table.df) == 3
    assert table.df.iloc[0].tolist() == ["北京", 1200, 300, 1500, 400]
    assert table.df["2024 年_销售额"].tolist() == [1500, 1400, 1250]


def test_title_row_plus_merged_header(tmp_path: Path):
    path = tmp_path / "report.xlsx"
    wb, ws = _new(path)
    ws["A1"] = "2024 年各城市销售报表"
    ws.merge_cells("A1:E1")
    _two_level_header(ws, start_row=2)
    wb.save(path)

    df = ExcelReader.read(str(path))[0].df

    assert list(df.columns) == [
        "城市",
        "2023 年_销售额",
        "2023 年_利润",
        "2024 年_销售额",
        "2024 年_利润",
    ]
    assert df["城市"].tolist() == ["北京", "上海", "广州"]


def test_title_row_with_normal_header(tmp_path: Path):
    path = tmp_path / "title_only.xlsx"
    wb, ws = _new(path)
    ws["A1"] = "2024 年销售报表"
    ws.merge_cells("A1:C1")
    ws.append(["城市", "销售额", "利润"])
    ws.append(["北京", 1500, 400])
    ws.append(["上海", 1400, 360])
    wb.save(path)

    df = ExcelReader.read(str(path))[0].df

    assert list(df.columns) == ["城市", "销售额", "利润"]
    assert df["销售额"].tolist() == [1500, 1400]


def test_normal_sheet_keeps_pandas_behaviour(tmp_path: Path):
    path = tmp_path / "normal.xlsx"
    wb, ws = _new(path)
    ws.append(["城市", "销售额"])
    ws.append(["北京", 1500])
    ws.append(["上海", 1400])
    wb.save(path)

    reader_df = ExcelReader.read(str(path))[0].df
    pandas_df = pd.read_excel(path)

    pd.testing.assert_frame_equal(reader_df, pandas_df)


def test_merged_data_cells_do_not_move_header(tmp_path: Path):
    path = tmp_path / "data_merge.xlsx"
    wb, ws = _new(path)
    ws.append(["城市", "销售额"])
    ws["A2"] = "北京"
    ws.merge_cells("A2:A3")
    ws["B2"] = 1500
    ws["B3"] = 1600
    wb.save(path)

    df = ExcelReader.read(str(path))[0].df

    assert list(df.columns) == ["城市", "销售额"]
    assert len(df) == 2
    assert df["销售额"].tolist() == [1500, 1600]


def test_multiple_sheets_and_sheet_name(tmp_path: Path):
    path = tmp_path / "multi.xlsx"
    wb = Workbook()
    summary = wb.active
    summary.title = "汇总"
    _two_level_header(summary)
    detail = wb.create_sheet("明细")
    detail.append(["月份", "销量"])
    detail.append(["1 月", 10])
    detail.append(["2 月", 20])
    wb.save(path)

    tables = ExcelReader.read(str(path))

    assert [t.name for t in tables] == ["multi", "multi_sheet2"]
    assert "2023 年_销售额" in tables[0].df.columns
    assert list(tables[1].df.columns) == ["月份", "销量"]

    by_name = ExcelReader.read(str(path), sheet_name="汇总")[0]
    assert "2024 年_利润" in by_name.df.columns

    by_index = ExcelReader.read(str(path), sheet_name=1)[0]
    assert list(by_index.df.columns) == ["月份", "销量"]


def test_duplicate_leaf_names_get_unique_suffix(tmp_path: Path):
    path = tmp_path / "dup.xlsx"
    wb, ws = _new(path)
    ws["A1"] = "城市"
    ws["B1"] = "销售额"
    ws["C1"] = "销售额"
    ws.merge_cells("A1:A2")
    ws.merge_cells("B1:C1")
    ws["B2"] = "2023"
    ws["C2"] = "2024"
    ws.append(["北京", 100, 200])
    wb.save(path)

    df = ExcelReader.read(str(path))[0].df

    # 纵向合并的 A 列只有一个名字；横向合并的子列名不同
    assert list(df.columns) == ["城市", "销售额_2023", "销售额_2024"]


@pytest.mark.parametrize("name", ["books.xlsx", "books.xlsm"])
def test_xlsm_and_xlsx_are_supported(tmp_path: Path, name: str):
    path = tmp_path / name
    wb, ws = _new(path)
    _two_level_header(ws)
    wb.save(path)

    df = ExcelReader.read(str(path))[0].df
    assert "2023 年_销售额" in df.columns


def test_blank_cell_two_row_header(tmp_path: Path):
    """没有合并且靠空白格表达的两行表头也要能拼列名。"""
    path = tmp_path / "unmerged.xlsx"
    wb, ws = _new(path)
    ws.append(["城市", "2023 年", None, "2024 年", None])
    ws.append([None, "销售额", "利润", "销售额", "利润"])
    ws.append(["北京", 1200, 300, 1500, 400])
    ws.append(["上海", 1100, 280, 1400, 360])
    wb.save(path)

    df = ExcelReader.read(str(path))[0].df

    assert list(df.columns) == [
        "城市",
        "2023 年_销售额",
        "2023 年_利润",
        "2024 年_销售额",
        "2024 年_利润",
    ]
    assert len(df) == 2
    assert df["2024 年_销售额"].tolist() == [1500, 1400]


def test_blank_cell_but_data_row_not_treated_as_header(tmp_path: Path):
    """第二行首列有值 + 含数字，是数据行，不能当表头吃掉。"""
    path = tmp_path / "not_header.xlsx"
    wb, ws = _new(path)
    ws.append(["城市", None, "地区"])
    ws.append(["北京", "华北", "北方"])
    ws.append(["上海", 1200, "南方"])
    wb.save(path)

    df = ExcelReader.read(str(path))[0].df

    assert list(df.columns) == ["城市", "Unnamed: 1", "地区"]
    assert df.iloc[0].tolist() == ["北京", "华北", "北方"]
    assert len(df) == 2


def test_blank_cell_with_numeric_second_row_is_data(tmp_path: Path):
    path = tmp_path / "numeric_row.xlsx"
    wb, ws = _new(path)
    ws.append(["城市", "销售额"])
    ws.append([None, 1200])
    ws.append(["上海", 1400])
    wb.save(path)

    df = ExcelReader.read(str(path))[0].df

    assert list(df.columns) == ["城市", "销售额"]
    assert len(df) == 2


async def test_merged_header_reaches_data_profile(tmp_path: Path):
    """走完整 ingest：合并表头列名要进入 schema 和 series，而不是 Unnamed。"""
    from service.data_ingestion import DataSource, ingest

    path = tmp_path / "merged_profile.xlsx"
    wb, ws = _new(path)
    _two_level_header(ws)
    wb.save(path)

    profile = await ingest(DataSource(kind="file", path=str(path), name="merged"))

    assert [c["name"] for c in profile.schema] == [
        "城市",
        "2023 年_销售额",
        "2023 年_利润",
        "2024 年_销售额",
        "2024 年_利润",
    ]
    assert profile.row_count == 3
    # 4 个「年份_指标」列是测量值，不再被误当成时间点
    assert profile.column_semantics["2023 年_销售额"] == "measure"
    assert len(profile.series_catalog.series) == 1
    series = profile.series_catalog.series[0]
    assert series.axis_kind == "categorical_comparison"
    assert series.row_label_column == "城市"
    assert series.column_label_columns == [
        "2023 年_销售额",
        "2023 年_利润",
        "2024 年_销售额",
        "2024 年_利润",
    ]


def test_xls_merged_ranges_are_converted(monkeypatch, tmp_path: Path):
    """xlrd 的 rhi/chi 是开区间，转成 1-based 闭区间后才能和 xlsx 逻辑共用。"""
    path = tmp_path / "old.xls"
    path.write_bytes(b"placeholder")

    class _Sheet:
        merged_cells = [(0, 2, 0, 1), (1, 2, 1, 3)]

    class _Book:
        def sheet_by_name(self, name):
            return _Sheet()

    import xlrd

    monkeypatch.setattr(xlrd, "open_workbook", lambda *args, **kwargs: _Book())

    assert _merged_ranges(str(path), "Sheet1") == [(1, 1, 2, 1), (2, 2, 2, 3)]


def test_xls_merged_ranges_falls_back_on_error(monkeypatch, tmp_path: Path):
    path = tmp_path / "broken.xls"
    path.write_bytes(b"not a real xls")

    import xlrd

    monkeypatch.setattr(
        xlrd, "open_workbook", lambda *args, **kwargs: (_ for _ in ()).throw(RuntimeError("bad"))
    )

    assert _merged_ranges(str(path), "Sheet1") is None


def test_merged_data_cells_fill_text_down(tmp_path: Path):
    """数据区纵向合并的文本（城市）要向下填满，不能让后续行变 NaN。"""
    path = tmp_path / "merged_data.xlsx"
    wb, ws = _new(path)
    ws.append(["城市", "月份", "销售额"])
    ws.append(["北京", "1月", 100])
    ws.append([None, "2月", 120])
    ws.merge_cells("A2:A3")
    ws.append(["上海", "1月", 90])
    ws.append([None, "2月", 95])
    ws.merge_cells("A4:A5")
    wb.save(path)

    df = ExcelReader.read(str(path))[0].df

    assert df["城市"].tolist() == ["北京", "北京", "上海", "上海"]
    assert df["月份"].tolist() == ["1月", "2月", "1月", "2月"]
    assert df["销售额"].tolist() == [100, 120, 90, 95]


def test_merged_numeric_total_not_filled_down(tmp_path: Path):
    """数值合并多为合计/小计，向下填会重复计量，必须保持 NaN。"""
    path = tmp_path / "merged_total.xlsx"
    wb, ws = _new(path)
    ws.append(["城市", "销售额", "备注"])
    ws.append(["北京", 100, None])
    ws.append([None, 200, None])
    ws.merge_cells("A2:A3")
    ws.append(["合计", 300, "小计"])
    ws.append([None, None, "尾部说明"])
    ws.merge_cells("B4:B5")
    wb.save(path)

    df = ExcelReader.read(str(path))[0].df

    assert df["城市"].tolist()[:2] == ["北京", "北京"]
    assert df.iloc[2]["城市"] == "合计"
    assert pd.isna(df.iloc[3]["销售额"])
    assert df.iloc[3]["备注"] == "尾部说明"


def test_merged_data_with_two_level_header(tmp_path: Path):
    """两行合并表头 + 数据区纵向合并：填充行号要按实际数据起始行对齐。"""
    path = tmp_path / "merged_both.xlsx"
    wb, ws = _new(path)
    ws["A1"] = "城市"
    ws.merge_cells("A1:A2")
    ws["B1"] = "2023 年"
    ws.merge_cells("B1:C1")
    ws["B2"] = "销售额"
    ws["C2"] = "利润"

    ws["A3"] = "北京"
    ws.merge_cells("A3:A4")
    ws["B3"], ws["C3"] = 100, 10
    ws["B4"], ws["C4"] = 120, 12
    ws["A5"], ws["B5"], ws["C5"] = "上海", 90, 9
    wb.save(path)

    df = ExcelReader.read(str(path))[0].df

    assert list(df.columns) == ["城市", "2023 年_销售额", "2023 年_利润"]
    assert df["城市"].tolist() == ["北京", "北京", "上海"]
    assert df["2023 年_销售额"].tolist() == [100, 120, 90]
