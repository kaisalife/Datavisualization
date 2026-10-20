"""表格生成效果评估（真实模型，needs_llm）。

评估两条链路：
- 文本 -> 表格：``LlmTabularReader``（Markdown 表 / 日志）
- 图片 -> 表格：``MultimodalReader``（表格截图 / 柱状图 / 无表格图）

运行（默认测试集不会跑）：

    .venv\\Scripts\\python.exe -m pytest tests/test_table_generation_eval.py \\
        -o addopts="" -m "needs_llm and needs_network" -q -s

模型配置来自 ``tests/.env``（见 tests/conftest.py）。
评分标准是“能否直接支撑画图”，而不是逐字一致：列名做模糊匹配，数值做容差比较。
"""
from __future__ import annotations

import re
from pathlib import Path

import pytest

pytestmark = [pytest.mark.needs_llm, pytest.mark.needs_network, pytest.mark.integration]

_FONT = r"C:\Windows\Fonts\msyh.ttc"
_NUM_RE = re.compile(r"-?\d+(?:\.\d+)?")


# ---------------------------------------------------------------- 评分工具


def _norm(text) -> str:
    """归一化：小写、去空白、全角括号转半角。"""
    s = str(text).strip().lower().replace("（", "(").replace("）", ")")
    return re.sub(r"\s+", "", s)


def _column_hits(expected: list[str], actual: list[str]) -> int:
    """expected 的每列能否在 actual 中找到（双向包含即算命中）。"""
    hits = 0
    for exp in expected:
        ne = _norm(exp)
        if any(ne in _norm(col) or _norm(col) in ne for col in actual):
            hits += 1
    return hits


def _key_column(df, key: str):
    """按列名包含关系找主键列，找不到退回第一列。"""
    nk = _norm(key)
    for col in df.columns:
        if nk in _norm(col) or _norm(col) in nk:
            return col
    return df.columns[0]


def _cell_number(value) -> float | None:
    """把单元格转成数字；无法转返回 None。"""
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        return float(value)
    match = _NUM_RE.search(str(value).replace(",", ""))
    return float(match.group()) if match else None


def _row_values(df, key_value: str, key: str) -> dict:
    """取主键匹配的那一行的 {列名: 原始值}。"""
    col = _key_column(df, key)
    target = _norm(key_value)
    matched = df[df[col].map(_norm) == target]
    if matched.empty:
        return {}
    return matched.iloc[0].to_dict()


def _all_numbers(df) -> list[float]:
    """收集表中所有数值单元格。"""
    values = []
    for col in df.columns:
        for cell in df[col].tolist():
            number = _cell_number(cell)
            if number is not None:
                values.append(number)
    return values


def _report(name: str, checks: dict[str, bool]) -> float:
    """打印单项评分并返回通过率。"""
    passed = sum(1 for ok in checks.values() if ok)
    total = len(checks)
    print(f"\n[{name}] {passed}/{total} " + " ".join(
        f"{k}={'PASS' if v else 'FAIL'}" for k, v in checks.items()
    ))
    return passed / total if total else 0.0


# ---------------------------------------------------------------- 造图


def _render_table(path: Path, headers: list[str], rows: list[list[str]]) -> None:
    """渲染带边框的中文表格截图。"""
    from PIL import Image, ImageDraw, ImageFont

    font = ImageFont.truetype(_FONT, 20)
    col_w, row_h, pad = 190, 52, 16
    width = col_w * len(headers) + pad * 2
    height = row_h * (len(rows) + 1) + pad * 2
    img = Image.new("RGB", (width, height), "white")
    draw = ImageDraw.Draw(img)

    for i in range(len(rows) + 2):
        y = pad + i * row_h
        draw.line([(pad, y), (width - pad, y)], fill=(60, 60, 60), width=2)
    for j in range(len(headers) + 1):
        x = pad + j * col_w
        draw.line([(x, pad), (x, pad + row_h * (len(rows) + 1))], fill=(60, 60, 60), width=2)

    for j, head in enumerate(headers):
        draw.text((pad + j * col_w + 14, pad + 14), head, fill="black", font=font)
    for i, row in enumerate(rows, start=1):
        for j, cell in enumerate(row):
            draw.text((pad + j * col_w + 14, pad + i * row_h + 14), str(cell), fill="black", font=font)

    img.save(path)


def _render_bar_chart(path: Path) -> None:
    """渲染无数据标签的柱状图（考察按坐标轴估读）。"""
    from PIL import Image, ImageDraw, ImageFont

    font = ImageFont.truetype(_FONT, 16)
    width, height = 640, 420
    left, bottom, top = 80, 360, 30
    img = Image.new("RGB", (width, height), "white")
    draw = ImageDraw.Draw(img)

    max_value = 250
    for tick in range(0, max_value + 1, 50):
        y = bottom - (tick / max_value) * (bottom - top)
        draw.line([(left, y), (width - 40, y)], fill=(220, 220, 220), width=1)
        draw.text((left - 45, y - 8), str(tick), fill="black", font=font)
    draw.line([(left, top), (left, bottom)], fill="black", width=2)
    draw.line([(left, bottom), (width - 40, bottom)], fill="black", width=2)

    labels = ["2023", "2024", "2025"]
    values = [100, 150, 200]
    slot = (width - 40 - left) / len(labels)
    bar_w = slot * 0.5
    for i, (label, value) in enumerate(zip(labels, values)):
        x0 = left + slot * i + (slot - bar_w) / 2
        y0 = bottom - (value / max_value) * (bottom - top)
        draw.rectangle([x0, y0, x0 + bar_w, bottom - 2], fill=(70, 130, 200))
        draw.text((x0 + bar_w / 2 - 18, bottom + 8), label, fill="black", font=font)

    draw.text((width / 2 - 40, 5), "年度销量", fill="black", font=font)
    img.save(path)


def _render_noise(path: Path) -> None:
    """渐变图：不含任何表格/数据，用于检验“不要臆造”。"""
    from PIL import Image

    img = Image.new("RGB", (480, 300), "white")
    pixels = img.load()
    for x in range(img.width):
        for y in range(img.height):
            pixels[x, y] = (200 - x // 4, 180 - y // 6, 160)
    img.save(path)


def _render_merged_table(path: Path) -> None:
    """城市列跨两行的合并单元格表格。"""
    from PIL import Image, ImageDraw, ImageFont

    font = ImageFont.truetype(_FONT, 20)
    headers = ["城市", "季度", "销量"]
    rows = [["北京", "Q1", "100"], ["", "Q2", "120"], ["上海", "Q1", "90"], ["", "Q2", "110"]]
    col_w, row_h, pad = 170, 50, 16
    img = Image.new("RGB", (col_w * 3 + pad * 2, row_h * 5 + pad * 2), "white")
    draw = ImageDraw.Draw(img)

    for i in range(len(rows) + 2):
        y = pad + i * row_h
        draw.line([(pad, y), (img.width - pad, y)], fill=(50, 50, 50), width=2)
    for j in range(len(headers) + 1):
        x = pad + j * col_w
        draw.line([(x, pad), (x, pad + row_h * 5)], fill=(50, 50, 50), width=2)

    for j, head in enumerate(headers):
        draw.text((pad + j * col_w + 12, pad + 13), head, fill="black", font=font)
    for i, row in enumerate(rows, start=1):
        for j, cell in enumerate(row):
            if cell:
                draw.text((pad + j * col_w + 12, pad + i * row_h + 13), cell, fill="black", font=font)

    # 擦掉合并区域内部的分隔线，只保留一个城市名
    for col, start, end, value in ((0, 1, 2, "北京"), (0, 3, 4, "上海")):
        x = pad + col * col_w
        y0, y1 = pad + start * row_h, pad + (end + 1) * row_h
        draw.rectangle([x + 2, y0 + 2, x + col_w - 2, y1 - 2], fill="white")
        draw.text((x + 12, (y0 + y1) // 2 - 12), value, fill="black", font=font)

    img.save(path)


def _render_multitable(path: Path) -> None:
    """左侧键值对（无关）+ 右侧 3x3 销售主表。"""
    from PIL import Image, ImageDraw, ImageFont

    font = ImageFont.truetype(_FONT, 20)
    img = Image.new("RGB", (900, 360), "white")
    draw = ImageDraw.Draw(img)
    draw.text((40, 30), "员工人数：12", fill="black", font=font)
    draw.text((40, 80), "部门数：3", fill="black", font=font)

    headers = ["城市", "2023年", "2024年"]
    rows = [["北京", "1200", "1500"], ["上海", "1100", "1400"], ["广州", "900", "1250"]]
    x0, y0, col_w, row_h = 340, 140, 170, 50
    draw.text((x0 + 120, 96), "区域销售明细", fill="black", font=font)
    for i in range(len(rows) + 2):
        draw.line([(x0, y0 + i * row_h), (x0 + col_w * 3, y0 + i * row_h)], fill=(50, 50, 50), width=2)
    for j in range(4):
        draw.line([(x0 + j * col_w, y0), (x0 + j * col_w, y0 + row_h * 4)], fill=(50, 50, 50), width=2)
    for j, head in enumerate(headers):
        draw.text((x0 + j * col_w + 12, y0 + 13), head, fill="black", font=font)
    for i, row in enumerate(rows, start=1):
        for j, cell in enumerate(row):
            draw.text((x0 + j * col_w + 12, y0 + i * row_h + 13), cell, fill="black", font=font)

    img.save(path)


def _render_two_independent_tables(path: Path) -> None:
    """左右两张列结构完全不同的表（销售 / 库存），检验一图多表。"""
    from PIL import Image, ImageDraw, ImageFont

    font = ImageFont.truetype(_FONT, 20)
    title_font = ImageFont.truetype(_FONT, 22)
    col_w, row_h, pad = 170, 50, 12

    def draw_table(draw, x0, y0, headers, rows):
        for i in range(len(rows) + 2):
            draw.line(
                [(x0, y0 + i * row_h), (x0 + col_w * len(headers), y0 + i * row_h)],
                fill=(50, 50, 50),
                width=2,
            )
        for j in range(len(headers) + 1):
            draw.line(
                [(x0 + j * col_w, y0), (x0 + j * col_w, y0 + row_h * (len(rows) + 1))],
                fill=(50, 50, 50),
                width=2,
            )
        for j, head in enumerate(headers):
            draw.text((x0 + j * col_w + 12, y0 + 13), head, fill="black", font=font)
        for i, row in enumerate(rows, start=1):
            for j, cell in enumerate(row):
                draw.text((x0 + j * col_w + 12, y0 + i * row_h + 13), cell, fill="black", font=font)

    width = col_w * 2 * 2 + pad * 3
    height = row_h * 4 + 90
    img = Image.new("RGB", (width, height), "white")
    draw = ImageDraw.Draw(img)

    left_x, right_x, top = pad, col_w * 2 + pad * 2, 52
    draw.text((left_x + 20, 14), "表 1：2024 年区域销售", fill="black", font=title_font)
    draw.text((right_x + 50, 14), "表 2：库存", fill="black", font=title_font)
    draw_table(draw, left_x, top, ["城市", "2024 年"], [["北京", "1500"], ["上海", "1400"], ["广州", "1250"]])
    draw_table(draw, right_x, top, ["产品", "库存"], [["A", "30"], ["B", "45"], ["C", "12"]])

    img.save(path)


# ---------------------------------------------------------------- 用例


def test_text_markdown_table(test_llm, tmp_path: Path):
    """Markdown 表格 -> 结构化表格：列名、行数、关键值。"""
    from service.data_ingestion.readers.llm_tabular_reader import LlmTabularReader

    md = tmp_path / "sales.md"
    md.write_text(
        "# 2024 年区域销售\n\n"
        "| 城市 | 销售额（万元） | 同比增长 |\n"
        "| --- | --- | --- |\n"
        "| 北京 | 1200 | 8.5 |\n"
        "| 上海 | 1100 | 6.2 |\n"
        "| 广州 | 900 | -1.5 |\n",
        encoding="utf-8",
    )

    df = LlmTabularReader.read(str(md), llm=test_llm)[0].df
    beijing = _row_values(df, "北京", "城市")
    guangzhou = _row_values(df, "广州", "城市")
    print("\n文本表 columns:", list(df.columns), "rows:", len(df))
    print(df.to_string(index=False))

    checks = {
        "列名>=2": _column_hits(["城市", "销售额", "同比增长"], list(df.columns)) >= 2,
        "行数=3": len(df) == 3,
        "北京销售额=1200": _cell_number(beijing.get("销售额", beijing.get(_key_column(df, "销售额"), None))) == 1200,
        "广州同比=-1.5": abs((_cell_number(
            guangzhou.get("同比增长", guangzhou.get(_key_column(df, "同比增长"), None)) or 0
        ) or 0) + 1.5) < 1e-6,
    }
    assert _report("文本-Markdown", checks) == 1.0


def test_text_log_records(test_llm, tmp_path: Path):
    """日志行 -> 表：识别重复字段，而不是逐条堆成一列。"""
    from service.data_ingestion.readers.llm_tabular_reader import LlmTabularReader

    log = tmp_path / "service.log"
    log.write_text(
        "2024-01-01 10:00:01 INFO  api  200  12ms\n"
        "2024-01-01 10:00:02 WARN  db   503  500ms\n"
        "2024-01-01 10:00:03 ERROR api  500  830ms\n"
        "2024-01-01 10:00:04 INFO  api  200  9ms\n",
        encoding="utf-8",
    )

    df = LlmTabularReader.read(str(log), llm=test_llm)[0].df
    print("\n日志表 columns:", list(df.columns), "rows:", len(df))
    print(df.to_string(index=False))

    checks = {
        "列数>=4": len(df.columns) >= 4,
        "行数=4": len(df) == 4,
        "含时间列": any("时间" in _norm(c) or "date" in _norm(c) for c in df.columns),
        "含状态码": any("状态" in _norm(c) or "status" in _norm(c) or "码" in _norm(c) for c in df.columns),
    }
    assert _report("文本-日志", checks) == 1.0


def test_image_table(test_llm, tmp_path: Path):
    """表格截图 -> 表格：中文列名、行数、数值原样还原。"""
    from service.data_ingestion.readers.multimodal_reader import MultimodalReader

    png = tmp_path / "table.png"
    _render_table(
        png,
        ["城市", "2023 年", "2024 年"],
        [["北京", "1200", "1500"], ["上海", "1100", "1400"], ["广州", "900", "1250"]],
    )

    df = MultimodalReader.read(str(png), llm=test_llm)[0].df
    beijing = _row_values(df, "北京", "城市")
    values_2024 = [
        _cell_number(v) for k, v in beijing.items()
        if "2024" in _norm(k)
    ]
    print("\n图片表 columns:", list(df.columns), "rows:", len(df))
    print(df.to_string(index=False))

    checks = {
        "列名与原图一致": [_norm(c) for c in df.columns] == ["城市", "2023年", "2024年"],
        "行数=3": len(df) == 3,
        "北京行存在": bool(beijing),
        "北京2024=1500": any(v == 1500 for v in values_2024 if v is not None),
    }
    assert _report("图片-表格", checks) == 1.0


def test_image_bar_chart_estimation(test_llm, tmp_path: Path):
    """柱状图截图 -> 按坐标轴估读数据，允许 12% 误差。"""
    from service.data_ingestion.readers.multimodal_reader import MultimodalReader

    png = tmp_path / "chart.png"
    _render_bar_chart(png)

    df = MultimodalReader.read(str(png), llm=test_llm)[0].df
    numbers = _all_numbers(df)
    expected = [100.0, 150.0, 200.0]
    matched = sum(
        1 for exp in expected
        if any(abs(num - exp) <= max(12.0, exp * 0.12) for num in numbers)
    )
    print("\n图表 columns:", list(df.columns), "rows:", len(df), "numbers:", numbers)
    print(df.to_string(index=False))

    checks = {
        "行数>=3": len(df) >= 3,
        "3 个值均命中(±12%)": matched == 3,
        "列名标注估算或数值": any(
            "估算" in _norm(c) or "值" in _norm(c) or "销量" in _norm(c) for c in df.columns
        ),
    }
    assert _report("图片-柱状图", checks) == 1.0


def test_image_without_table_does_not_hallucinate(test_llm, tmp_path: Path):
    """无表格的图片应拒绝输出，而不是编造数据。"""
    from service.data_ingestion.readers.multimodal_reader import MultimodalReader

    png = tmp_path / "noise.png"
    _render_noise(png)

    with pytest.raises(ValueError):
        MultimodalReader.read(str(png), llm=test_llm)

    print("\n[图片-无表格] PASS（未臆造）")


def test_image_merged_cells(test_llm, tmp_path: Path):
    """合并单元格：城市列跨两行，应向下填充而不是留空。"""
    from service.data_ingestion.readers.multimodal_reader import MultimodalReader

    png = tmp_path / "merged.png"
    _render_merged_table(png)

    df = MultimodalReader.read(str(png), llm=test_llm)[0].df
    city_col = _key_column(df, "城市")
    counts = {_norm(k): int(v) for k, v in df[city_col].value_counts().items()}
    amount_col = _key_column(df, "销")
    total = sum(_cell_number(v) or 0 for v in df[amount_col].tolist())
    print("\n合并单元格 columns:", list(df.columns), "rows:", len(df))
    print(df.to_string(index=False))

    checks = {
        "行数=4": len(df) == 4,
        "北京填充2行": counts.get("北京") == 2,
        "上海填充2行": counts.get("上海") == 2,
        "销量合计=420": total == 420,
    }
    assert _report("图片-合并单元格", checks) == 1.0


def test_image_multiple_tables_picks_main(test_llm, tmp_path: Path):
    """同图多表：应选数据最完整的主表，忽略旁边的键值对。"""
    from service.data_ingestion.readers.multimodal_reader import MultimodalReader

    png = tmp_path / "multi.png"
    _render_multitable(png)

    df = MultimodalReader.read(str(png), llm=test_llm)[0].df
    beijing = _row_values(df, "北京", "城市")
    values_2024 = [_cell_number(v) for k, v in beijing.items() if "2024" in _norm(k)]
    print("\n多表 columns:", list(df.columns), "rows:", len(df))
    print(df.to_string(index=False))

    checks = {
        "行数=3（选中主表）": len(df) == 3,
        "含城市列": any("城市" in _norm(c) for c in df.columns),
        "含2023/2024列": _column_hits(["2023", "2024"], list(df.columns)) == 2,
        "北京2024=1500": any(v == 1500 for v in values_2024 if v is not None),
    }
    assert _report("图片-多表", checks) == 1.0


def test_docx_embedded_table_image(test_llm, tmp_path: Path):
    """Word 里贴的表格截图 -> 表格（走 OOXML 内嵌图片链路）。"""
    from docx import Document
    from docx.shared import Inches

    from service.data_ingestion.readers.office_image_reader import OfficeImageReader

    png = tmp_path / "table.png"
    _render_table(
        png,
        ["城市", "2023 年", "2024 年"],
        [["北京", "1200", "1500"], ["上海", "1100", "1400"], ["广州", "900", "1250"]],
    )

    doc = Document()
    doc.add_paragraph("附件：2024 年区域销售情况（截图）")
    doc.add_picture(str(png), width=Inches(6))
    docx_path = tmp_path / "sales_report.docx"
    doc.save(str(docx_path))

    tables = OfficeImageReader.read(str(docx_path), table_name="区域销售", llm=test_llm)
    df = tables[0].df
    beijing = _row_values(df, "北京", "城市")
    values_2024 = [_cell_number(v) for k, v in beijing.items() if "2024" in _norm(k)]
    print("\nDOCX内嵌图 columns:", list(df.columns), "rows:", len(df), "path:", tables[0].source_path)
    print(df.to_string(index=False))

    checks = {
        "只抽到1张表": len(tables) == 1,
        "列名与原图一致": [_norm(c) for c in df.columns] == ["城市", "2023年", "2024年"],
        "行数=3": len(df) == 3,
        "北京2024=1500": any(v == 1500 for v in values_2024 if v is not None),
    }
    assert _report("DOCX-内嵌表格截图", checks) == 1.0


def test_image_multi_table_extraction(test_llm, tmp_path: Path):
    """一图两表：multi_table=True 时两张都应抽出，且不互相污染。"""
    from service.data_ingestion.readers.multimodal_reader import MultimodalReader

    png = tmp_path / "two_independent_tables.png"
    _render_two_independent_tables(png)

    tables = MultimodalReader.read(str(png), llm=test_llm, multi_table=True)
    for t in tables:
        print(f"\n[{t.name}] columns={list(t.df.columns)} rows={len(t.df)}")
        print(t.df.to_string(index=False))

    sales = next(
        (t.df for t in tables if any("城市" in _norm(c) for c in t.df.columns)), None
    )
    stock = next(
        (t.df for t in tables if any("产品" in _norm(c) for c in t.df.columns)), None
    )

    beijing = _row_values(sales, "北京", "城市") if sales is not None else {}
    values_2024 = [_cell_number(v) for k, v in beijing.items() if "2024" in _norm(k)]
    a_row = _row_values(stock, "A", "产品") if stock is not None else {}
    a_stock = (
        _cell_number(a_row.get("库存", a_row.get(_key_column(stock, "库存"))))
        if a_row
        else None
    )

    checks = {
        "至少2张表": len(tables) >= 2,
        "销售表3行": sales is not None and len(sales) == 3,
        "北京2024=1500": any(v == 1500 for v in values_2024 if v is not None),
        "库存表3行": stock is not None and len(stock) == 3,
        "库存A=30": a_stock == 30,
    }
    assert _report("图片-一图多表", checks) == 1.0
